# Copyright 2026 Evan Montgomery-Recht
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.
"""Build the progress-tracker workbook: By section, By lane, Chapters, About."""

from __future__ import annotations

import datetime as _dt
from dataclasses import dataclass

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .catalog import Duolingo, Textbook
from .lanes import chapter_lanes
from .mapping import RATING_ORDER, Mapping

STATUSES = ("Not started", "In progress", "Done")
STATUS_FILL = {"Not started": "F2F2F2", "In progress": "FFF2CC", "Done": "C6EFCE"}
FIXED_TIME = _dt.datetime(2026, 1, 1, 0, 0, 0)

SECTION_COLUMNS = ("Chapter", "Section", "Section title", "Page", "Item", "Lane", "Unit #", "Unit", "Rating", "Status", "Date done", "Notes")
LANE_COLUMNS = ("Chapter", "Lane", "Unit #", "Unit", "Rating", "Sections", "Status", "Date done", "Notes")
SECTION_STATUS_COL = SECTION_COLUMNS.index("Status") + 1  # 1-based
LANE_STATUS_COL = LANE_COLUMNS.index("Status") + 1


@dataclass(frozen=True)
class SectionRow:
  chapter: str
  section: str
  title: str
  page: str
  item: str  # "Read section" or "Duolingo unit"
  lane: str
  unit_order: int | None
  unit: str
  rating: str


@dataclass(frozen=True)
class LaneRow:
  chapter: str
  lane: str
  unit_order: int
  unit: str
  rating: str
  sections: str


def _chapter_cell(chapter_id: str):
  return int(chapter_id) if chapter_id.isdigit() else chapter_id


def section_rows(book: Textbook, course: Duolingo, mapping: Mapping) -> list[SectionRow]:
  rows: list[SectionRow] = []
  for ch in book.chapters:
    for s in ch.sections:
      rows.append(SectionRow(ch.id, s.id, s.title, s.page, "Read section", "", None, "", ""))
      entry = mapping.entry(s.id)
      if not entry:
        continue
      refs = sorted(entry.units, key=lambda r: (RATING_ORDER[r.rating], course.unit(r.unit).lane, course.unit(r.unit).order))
      for ref in refs:
        u = course.unit(ref.unit)
        rows.append(SectionRow(ch.id, s.id, s.title, s.page, "Duolingo unit", course.lane(u.lane).name, u.order, u.name, ref.rating))
  return rows


def lane_rows(book: Textbook, course: Duolingo, mapping: Mapping) -> list[LaneRow]:
  rows: list[LaneRow] = []
  for ch in book.chapters:
    for lr in chapter_lanes(ch, mapping, course):
      for lu in lr.units:
        rows.append(LaneRow(ch.id, lr.lane.name, lu.unit.order, lu.unit.name, lu.rating, ", ".join(lu.sections)))
  return rows


def _write_sheet(ws, columns, rows, status_col: int, widths: dict[str, int]):
  ws.append(list(columns))
  for cell in ws[1]:
    cell.font = Font(bold=True)
    cell.alignment = Alignment(vertical="center")
  for r in rows:
    ws.append(r)
    for cell in ws[ws.max_row]:
      if isinstance(cell.value, str) and cell.value.startswith("="):
        cell.data_type = "s"  # catalog text is never a formula
  n = max(len(rows), 1) + 1
  for i, name in enumerate(columns, 1):
    ws.column_dimensions[get_column_letter(i)].width = widths.get(name, 14)
  ws.freeze_panes = "A2"
  ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{n}"
  col = get_column_letter(status_col)
  dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=False, showDropDown=False)
  dv.error = "Choose Not started, In progress, or Done."
  dv.errorTitle = "Status"
  ws.add_data_validation(dv)
  dv.add(f"{col}2:{col}{n}")
  for status, color in STATUS_FILL.items():
    ws.conditional_formatting.add(f"{col}2:{col}{n}", CellIsRule(operator="equal", formula=[f'"{status}"'],
                                                                 fill=PatternFill(start_color=color, end_color=color, fill_type="solid")))


def build_workbook(book: Textbook, course: Duolingo, mapping: Mapping, mapping_version: str = "unversioned") -> Workbook:
  wb = Workbook()
  wb.properties.creator = "precalc-tutor build"
  wb.properties.lastModifiedBy = "precalc-tutor build"
  wb.properties.created = FIXED_TIME
  wb.properties.modified = FIXED_TIME

  ws = wb.active
  ws.title = "By section"
  srows = [(_chapter_cell(r.chapter), r.section, r.title, r.page, r.item, r.lane, r.unit_order, r.unit, r.rating, STATUSES[0], None, None)
           for r in section_rows(book, course, mapping)]
  _write_sheet(ws, SECTION_COLUMNS, srows, SECTION_STATUS_COL,
               {"Chapter": 9, "Section": 9, "Section title": 40, "Page": 7, "Item": 14, "Lane": 18, "Unit #": 7, "Unit": 46, "Rating": 12, "Status": 13, "Date done": 12, "Notes": 30})

  wl = wb.create_sheet("By lane")
  lrows = [(_chapter_cell(r.chapter), r.lane, r.unit_order, r.unit, r.rating, r.sections, STATUSES[0], None, None)
           for r in lane_rows(book, course, mapping)]
  _write_sheet(wl, LANE_COLUMNS, lrows, LANE_STATUS_COL,
               {"Chapter": 9, "Lane": 18, "Unit #": 7, "Unit": 46, "Rating": 12, "Sections": 14, "Status": 13, "Date done": 12, "Notes": 30})

  wc = wb.create_sheet("Chapters")
  wc.append(["Chapter", "Title", "Items", "Done", "Percent"])
  for cell in wc[1]:
    cell.font = Font(bold=True)
  status_letter = get_column_letter(SECTION_STATUS_COL)
  for i, ch in enumerate(book.chapters, 2):
    wc.append([_chapter_cell(ch.id), ch.title,
               f"=COUNTIF('By section'!$A:$A,A{i})",
               f"=COUNTIFS('By section'!$A:$A,A{i},'By section'!${status_letter}:${status_letter},\"Done\")",
               f"=IF(C{i}=0,0,D{i}/C{i})"])
    wc.cell(row=i, column=5).number_format = "0%"
  wc.column_dimensions["B"].width = 44
  wc.freeze_panes = "A2"

  wa = wb.create_sheet("About")
  about = [
    ["Progress tracker", f"{book.title}, {book.edition}e, with Duolingo Math"],
    ["", ""],
    ["How to use", "Mark each row's Status as you go. 'By section' follows the book; 'By lane' lists the same units in the order Duolingo presents them, one block per grade or topic. The Chapters sheet totals the 'By section' sheet."],
    ["Status values", " / ".join(STATUSES)],
    ["Rating", "full: the unit covers the section's main skill. partial: covers part of it. prerequisite: a warm-up from an earlier grade."],
    ["", ""],
    ["Textbook", book.citation()],
    ["Duolingo inventory", f"{course.title}, {len(course.units())} units, last verified {course.last_verified} on {course.platform}" + ("" if course.complete else " (INCOMPLETE: some units may be missing)")],
    ["Mapping version", mapping_version],
    ["Mapped chapters", ", ".join(c.id for c in book.chapters if any(mapping.entry(s.id) and mapping.entry(s.id).units for s in c.sections)) or "none"],
    ["", ""],
    ["License", "Data and this workbook: CC BY 4.0. Textbook titles and Duolingo unit names are third-party material cited for reference; see the repository NOTICE."],
  ]
  for row in about:
    wa.append(row)
  wa.column_dimensions["A"].width = 20
  wa.column_dimensions["B"].width = 110
  for r in wa.iter_rows():
    r[0].font = Font(bold=True)
    r[1].alignment = Alignment(wrap_text=True, vertical="top")
  return wb
