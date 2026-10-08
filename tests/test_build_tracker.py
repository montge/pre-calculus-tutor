# Copyright 2026 Evan Montgomery-Recht
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.
import io
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from precalc_tutor import MAPPING_PATH
from precalc_tutor.catalog import load_duolingo, load_textbook
from precalc_tutor.checklist_pdf import build_pdf
from precalc_tutor.mapping import load_mapping
from precalc_tutor.tracker import STATUSES, build_workbook, lane_rows, section_rows
from tests.helpers import modified, write_temp_data

ROOT = Path(__file__).resolve().parent.parent
SCRIPT = ROOT / "scripts" / "build_tracker.py"


def _wb_bytes(wb) -> bytes:
  buf = io.BytesIO()
  wb.save(buf)
  return buf.getvalue()


def _cells(wb):
  out = {}
  for ws in wb:
    out[ws.title] = [[c.value for c in row] for row in ws.iter_rows()]
    out[ws.title + ":dv"] = sorted((dv.formula1, str(dv.sqref)) for dv in ws.data_validations.dataValidation)
  return out


class BuildTrackerTest(unittest.TestCase):
  @classmethod
  def setUpClass(cls):
    cls.book, cls.course, cls.mapping = load_textbook(), load_duolingo(), load_mapping()
    cls.wb = load_workbook(io.BytesIO(_wb_bytes(build_workbook(cls.book, cls.course, cls.mapping, "test"))))

  def test_section_rows_shape(self):
    rows = section_rows(self.book, self.course, self.mapping)
    r17 = [r for r in rows if r.section == "1.7"]
    self.assertEqual(r17[0].item, "Read section")
    self.assertTrue(all(r.item == "Duolingo unit" for r in r17[1:]))
    self.assertEqual(len(r17), 1 + len(self.mapping.entry("1.7").units))
    r23 = [r for r in rows if r.section == "2.3"]
    self.assertEqual(len(r23), 1)
    self.assertEqual(r23[0].item, "Read section")
    self.assertEqual(r23[0].unit, "")
    # consecutive: the 1.7 rows occupy one contiguous slice
    idx = [i for i, r in enumerate(rows) if r.section == "1.7"]
    self.assertEqual(idx, list(range(idx[0], idx[0] + len(idx))))

  def test_lane_rows_in_app_order_and_once_per_chapter(self):
    rows = lane_rows(self.book, self.course, self.mapping)
    ch3_g11 = [r for r in rows if r.chapter == "3" and r.lane == "Grade 11"]
    orders = [r.unit_order for r in ch3_g11]
    self.assertEqual(orders, sorted(orders))
    growth = [r for r in ch3_g11 if r.unit == "Exponential growth"]
    self.assertEqual(len(growth), 1)
    self.assertEqual(growth[0].sections, "3.1, 3.5")
    self.assertFalse(any(r.chapter == "4" for r in rows))

  def test_formula_like_text_is_stored_as_text(self):
    from openpyxl import Workbook
    from precalc_tutor.tracker import _write_sheet
    wb = Workbook(); ws = wb.active
    _write_sheet(ws, ("A", "Status"), [("=1+1", STATUSES[0])], 2, {})
    self.assertEqual(ws["A2"].data_type, "s")
    self.assertEqual(ws["A2"].value, "=1+1")

  def test_sheets_and_status_validation(self):
    self.assertEqual(self.wb.sheetnames, ["By section", "By lane", "Chapters", "About"])
    for name, col in (("By section", "J"), ("By lane", "G")):
      ws = self.wb[name]
      dvs = ws.data_validations.dataValidation
      self.assertEqual(len(dvs), 1, name)
      self.assertEqual(dvs[0].formula1, '"Not started,In progress,Done"')
      self.assertEqual(str(dvs[0].sqref), f"{col}2:{col}{ws.max_row}")
      self.assertEqual(ws[f"{col}1"].value, "Status")
      self.assertTrue(all(ws[f"{col}{r}"].value == STATUSES[0] for r in range(2, ws.max_row + 1)))
      self.assertEqual(ws.freeze_panes, "A2")
      self.assertTrue(ws.auto_filter.ref)

  def test_chapters_formulas(self):
    ws = self.wb["Chapters"]
    self.assertEqual(ws["C2"].value, "=COUNTIF('By section'!$A:$A,A2)")
    self.assertEqual(ws["D2"].value, "=COUNTIFS('By section'!$A:$A,A2,'By section'!$J:$J,\"Done\")")
    self.assertEqual(ws["E2"].value, "=IF(C2=0,0,D2/C2)")
    self.assertEqual(ws.max_row, 1 + len(self.book.chapters))

  def test_about_sheet(self):
    text = " ".join(str(c.value) for row in self.wb["About"].iter_rows() for c in row if c.value)
    self.assertIn("Precalculus with Limits", text)
    self.assertIn("Larson", text)
    self.assertIn(self.course.last_verified, text)
    self.assertIn("Mapped chapters 1, 2, 3", text)

  def test_workbook_reproducible(self):
    a = load_workbook(io.BytesIO(_wb_bytes(build_workbook(self.book, self.course, self.mapping, "v"))))
    b = load_workbook(io.BytesIO(_wb_bytes(build_workbook(self.book, self.course, self.mapping, "v"))))
    self.assertEqual(_cells(a), _cells(b))
    self.assertEqual(a.properties.created, b.properties.created)

  def test_pdf_reproducible_and_has_chapters(self):
    one = build_pdf(self.book, self.course, self.mapping)
    two = build_pdf(self.book, self.course, self.mapping)
    self.assertEqual(one, two)
    self.assertTrue(one.startswith(b"%PDF"))
    try:
      from pypdf import PdfReader
    except ImportError:
      self.skipTest("pypdf not installed")
    text = "".join(p.extract_text() for p in PdfReader(io.BytesIO(one)).pages)
    for title in ("Functions and Their Graphs", "Polynomial and Rational Functions", "Exponential and Logarithmic Functions"):
      self.assertIn(title, text)
    self.assertIn("Grade 11", text)
    self.assertIn("Book only, not on Duolingo", text)
    self.assertNotIn("Trigonometry", text)

  def test_script_refuses_invalid_data(self):
    def mutate(d):
      d["sections"][0]["units"][0]["unit"] = "u-g99-01"
    data = write_temp_data(mapping=modified(MAPPING_PATH, mutate))
    out = Path(tempfile.mkdtemp())
    proc = subprocess.run([sys.executable, str(SCRIPT), "--mapping", str(data / "mappings" / MAPPING_PATH.name),
                           "--textbook", str(data / "textbook" / "larson-precalculus-with-limits-2e.yaml"),
                           "--duolingo", str(data / "duolingo" / "math-course.yaml"),
                           "--xlsx", str(out / "t.xlsx"), "--pdf", str(out / "t.pdf")], capture_output=True, text=True, cwd=ROOT)
    self.assertEqual(proc.returncode, 1)
    self.assertIn("unknown unit", proc.stderr)
    self.assertFalse((out / "t.xlsx").exists())
    self.assertFalse((out / "t.pdf").exists())

  def test_committed_outputs_are_current(self):
    committed = load_workbook(ROOT / "dist" / "progress-tracker.xlsx")
    fresh = load_workbook(io.BytesIO(_wb_bytes(build_workbook(self.book, self.course, self.mapping, committed["About"]["B9"].value))))
    self.assertEqual(_cells(committed), _cells(fresh))
    self.assertEqual((ROOT / "dist" / "progress-checklist.pdf").read_bytes(), build_pdf(self.book, self.course, self.mapping))

  @unittest.skipUnless(shutil.which("soffice"), "LibreOffice not installed")
  def test_libreoffice_recalculates_zero_done(self):
    out = Path(tempfile.mkdtemp())
    src = out / "in.xlsx"
    src.write_bytes(_wb_bytes(build_workbook(self.book, self.course, self.mapping, "v")))
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(out / "lo"), str(src)],
                   capture_output=True, timeout=180, check=True)
    ws = load_workbook(out / "lo" / "in.xlsx", data_only=True)["Chapters"]
    self.assertGreater(ws["C2"].value, 0)
    self.assertEqual(ws["D2"].value, 0)
    self.assertEqual(ws["E2"].value, 0)


if __name__ == "__main__":
  unittest.main()
